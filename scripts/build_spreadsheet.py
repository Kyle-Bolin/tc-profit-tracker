"""Build spreadsheet/tc-profit-calculator.xlsx.

The workbook has four tabs:
  How to use  - short guide
  Settings    - global defaults (blue cells are inputs)
  Calculator  - one row per plant: inputs, then cost and profit outputs
  Price Log   - one row per observed listing or sale

Starting plant values are rough estimates, not checked against live listings.
Re-run this script to regenerate the template; it overwrites the file.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parent.parent / "spreadsheet" / "tc-profit-calculator.xlsx"

INPUT = Font(color="1F4E9E")
BOLD = Font(bold=True)
HEAD_FILL = PatternFill("solid", fgColor="E8EEF7")
OUT_FILL = PatternFill("solid", fgColor="F3F3F3")
WRAP = Alignment(wrap_text=True, vertical="top")

# name, retail (small rooted), wholesale plug, true-to-type risk,
# multiplication per cycle, cycle weeks, monthly price decay
PLANTS = [
    ("Monstera 'Thai Constellation'", 50, 5, "Low", 3, 6, 0.02),
    ("Philodendron 'Pink Princess'", 25, 3, "Medium", 3, 6, 0.02),
    ("Anthurium warocqueanum", 120, 18, "Low", 2.5, 8, 0.01),
    ("Philodendron 'Spiritus Sancti'", 180, 25, "Low", 2.5, 8, 0.04),
    ("Anthurium crystallinum", 35, 4, "Low", 3, 6, 0.01),
    ("Anthurium clarinervium", 40, 5, "Low", 3, 6, 0.01),
    ("Philodendron gloriosum", 40, 6, "Low", 2.5, 8, 0.01),
    ("Philodendron melanochrysum", 45, 6, "Low", 3, 6, 0.015),
    ("Philodendron 'Florida Beauty'", 60, 8, "Medium", 3, 6, 0.02),
    ("Philodendron 'White Princess'", 30, 4, "Medium", 3, 6, 0.02),
    ("Monstera 'Aurea' variegata", 150, 20, "Medium", 3, 6, 0.03),
    ("Monstera obliqua 'Peru'", 80, 10, "Low", 3, 6, 0.02),
    ("Alocasia 'Dragon Scale'", 20, 2, "Low", 4, 5, 0.01),
    ("Alocasia 'Black Velvet'", 20, 2, "Low", 4, 5, 0.01),
    ("Alocasia 'Frydek' variegata", 60, 8, "Medium", 3.5, 5, 0.03),
    ("Syngonium 'Albo'", 30, 3, "High", 4, 5, 0.02),
    ("Rhaphidophora tetrasperma", 15, 1.5, "Low", 4, 5, 0.01),
    ("Scindapsus treubii 'Moonlight'", 15, 1.5, "Low", 4, 5, 0.01),
    ("Hoya carnosa 'Krimson Queen'", 15, 2, "Medium", 3, 6, 0.005),
    ("Hoya kerrii variegata", 25, 3, "Medium", 3, 6, 0.01),
]

SETTINGS = [
    # key, label, value, format, note
    ("labor", "Labor rate ($/hour)", 15, "$#,##0.00", "What your time is worth. Set to 0 to see cash cost only."),
    ("tph", "Transfers per hour", 60, "0", "Explants cut and moved per hour. Commercial techs do far more."),
    ("per_vessel", "Explants per vessel", 5, "0", "How many plantlets share one jar."),
    ("vessel_cost", "Media + consumables per vessel ($)", 0.40, "$#,##0.00", "Agar, sugar, salts, hormones, PPM, alcohol, jar wear."),
    ("contam", "Contamination loss per cycle", 0.10, "0%", "Hobby range is often 5 to 25%."),
    ("acclim", "Rooting and acclimatization loss", 0.20, "0%", "Plantlets lost after leaving the jar."),
    ("acclim_months", "Months from deflask to sellable", 3, "0.0", "Humidity dome and growing-on time."),
    ("init_extra", "Initiation cost on top of mother plant ($)", 100, "$#,##0", "Failed sterilization attempts, supplies, time."),
    ("overhead", "Overhead per month ($)", 30, "$#,##0", "Lights, power, shelving and gear spread over time."),
    ("pack", "Packaging per plant sold ($)", 1.5, "$#,##0.00", "Box, pot, heat pack. Assumes buyer pays postage."),
    ("fee", "Platform fees", 0.13, "0%", "eBay or Etsy take including payment processing."),
    ("target", "Target sellable plants", 100, "0", "Batch size for the cost per plant."),
]


def build():
    wb = Workbook()

    # How to use
    ws = wb.active
    ws.title = "How to use"
    lines = [
        "Tissue Culture Profit Calculator",
        "",
        "Blue numbers are inputs you can change. Grey columns are calculated.",
        "Settings holds defaults used by every plant. On the Calculator tab, the per-plant columns for initiation, contamination, acclimatization loss and target pull from Settings until you type over them.",
        "Retail, plug prices, multiplication rates and decay rates are rough starting estimates, not checked against current listings. Replace them as the Price Log fills in.",
        "",
        "How the model works",
        "1. Plantlets to deflask = target / (1 - acclimatization loss).",
        "2. Each cycle multiplies plantlets by (multiplication rate x (1 - contamination)). Cycles needed = enough to reach step 1 from one clean explant.",
        "3. Plantlets handled = every transfer across all cycles plus the final deflask. Labor and media scale with this.",
        "4. Months to sellable = cycles x cycle weeks, plus growing-on time.",
        "5. Sale price = today's retail reduced by the monthly price decay for that many months, then minus platform fees.",
        "6. Profit per plant = net sale price - all-in cost per sellable plant.",
        "",
        "Price Log",
        "Add one row per listing or sale you see. Use the same plant name as the Calculator. Status 'sold' rows feed the Sold last 30 days and Avg sold price columns.",
        "Keep size classes separate: a TC plug and a 4-inch pot are different products.",
        "Stability rating is set by hand for now: Steady, Rising, Hype, Saturating, or Unknown.",
        "",
        "Legal reminder: many named cultivars are patented or protected. Growing for yourself is one thing; selling clones may need a license. Not legal advice.",
    ]
    for i, text in enumerate(lines, 1):
        ws.cell(i, 1, text).alignment = WRAP
    ws["A1"].font = Font(bold=True, size=14)
    ws["A7"].font = BOLD
    ws["A15"].font = BOLD
    ws.column_dimensions["A"].width = 110

    # Settings
    st = wb.create_sheet("Settings")
    st.append(["Setting", "Value", "Notes"])
    for c in st[1]:
        c.font = BOLD
        c.fill = HEAD_FILL
    ref = {}
    for i, (key, label, value, fmt, note) in enumerate(SETTINGS, 2):
        st.cell(i, 1, label)
        v = st.cell(i, 2, value)
        v.font = INPUT
        v.number_format = fmt
        st.cell(i, 3, note)
        ref[key] = f"Settings!$B${i}"
    st.column_dimensions["A"].width = 42
    st.column_dimensions["B"].width = 12
    st.column_dimensions["C"].width = 70

    # Price Log
    pl = wb.create_sheet("Price Log")
    pl_head = ["Date", "Plant", "Source", "Size class", "Price ($)", "Status", "Link", "Notes"]
    pl.append(pl_head)
    for c in pl[1]:
        c.font = BOLD
        c.fill = HEAD_FILL
    for col, w in zip("ABCDEFGH", [12, 34, 14, 16, 10, 10, 40, 40]):
        pl.column_dimensions[col].width = w
    for r in range(2, 2001):
        pl.cell(r, 1).number_format = "yyyy-mm-dd"
        pl.cell(r, 5).number_format = "$#,##0.00"
    dv_size = DataValidation(type="list", formula1='"cutting,TC plug,flask,small rooted,medium,large"', allow_blank=True)
    dv_status = DataValidation(type="list", formula1='"sold,listed"', allow_blank=True)
    dv_source = DataValidation(type="list", formula1='"eBay,Etsy,Palmstreet,Facebook,Shop,Wholesale,Other"', allow_blank=True)
    for dv, col in ((dv_size, "D"), (dv_status, "F"), (dv_source, "C")):
        pl.add_data_validation(dv)
        dv.add(f"{col}2:{col}2000")
    pl.freeze_panes = "A2"

    # Calculator
    ca = wb.create_sheet("Calculator", 1)
    head = [
        ("Plant", 32, None),
        ("Retail today ($)", 11, "$#,##0"),
        ("Wholesale plug ($)", 11, "$#,##0.00"),
        ("True-to-type risk", 10, None),
        ("Multiplication per cycle", 12, "0.0"),
        ("Cycle length (weeks)", 10, "0"),
        ("Price decay per month", 10, "0.0%"),
        ("Initiation cost ($)", 11, "$#,##0"),
        ("Contamination loss", 11, "0%"),
        ("Acclim. loss", 10, "0%"),
        ("Target sellable", 10, "0"),
        ("Plantlets to deflask", 11, "0"),
        ("Cycles needed", 9, "0"),
        ("Months to sellable", 10, "0.0"),
        ("Plantlets handled", 11, "0"),
        ("Labor hours", 9, "0.0"),
        ("Labor cost ($)", 11, "$#,##0"),
        ("Media cost ($)", 10, "$#,##0"),
        ("Overhead ($)", 10, "$#,##0"),
        ("Packaging ($)", 10, "$#,##0"),
        ("Total cost ($)", 11, "$#,##0"),
        ("Cost per sellable ($)", 11, "$#,##0.00"),
        ("Expected sale price ($)", 11, "$#,##0.00"),
        ("Net after fees ($)", 11, "$#,##0.00"),
        ("Profit per plant ($)", 11, "$#,##0.00"),
        ("Batch profit ($)", 11, "$#,##0"),
        ("Return on cost", 9, "0%"),
        ("Sold last 30 days", 9, "0"),
        ("Avg sold price, 90 days ($)", 11, "$#,##0.00"),
        ("Stability rating", 12, None),
    ]
    for j, (name, width, _) in enumerate(head, 1):
        c = ca.cell(1, j, name)
        c.font = BOLD
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="bottom")
        ca.column_dimensions[c.column_letter].width = width
    ca.row_dimensions[1].height = 45

    for i, (name, retail, plug, risk, mult, weeks, decay) in enumerate(PLANTS, 2):
        r = str(i)
        inputs = {
            "A": name, "B": retail, "C": plug, "D": risk, "E": mult, "F": weeks, "G": decay,
            "H": f"=B{r}+{ref['init_extra']}",
            "I": f"={ref['contam']}",
            "J": f"={ref['acclim']}",
            "K": f"={ref['target']}",
        }
        outputs = {
            "L": f"=K{r}/(1-J{r})",
            "M": f"=MAX(1,ROUNDUP(LN(L{r})/LN(E{r}*(1-I{r})),0))",
            "N": f"=M{r}*F{r}/4.345+{ref['acclim_months']}",
            "O": f"=L{r}/(1-1/(E{r}*(1-I{r})))+L{r}",
            "P": f"=O{r}/{ref['tph']}",
            "Q": f"=P{r}*{ref['labor']}",
            "R": f"=O{r}/{ref['per_vessel']}*{ref['vessel_cost']}",
            "S": f"=N{r}*{ref['overhead']}",
            "T": f"=K{r}*{ref['pack']}",
            "U": f"=H{r}+Q{r}+R{r}+S{r}+T{r}",
            "V": f"=U{r}/K{r}",
            "W": f"=B{r}*(1-G{r})^N{r}",
            "X": f"=W{r}*(1-{ref['fee']})",
            "Y": f"=X{r}-V{r}",
            "Z": f"=Y{r}*K{r}",
            "AA": f"=Z{r}/U{r}",
            "AB": f"=COUNTIFS('Price Log'!$B:$B,A{r},'Price Log'!$F:$F,\"sold\",'Price Log'!$A:$A,\">=\"&(TODAY()-30))",
            "AC": f"=IFERROR(AVERAGEIFS('Price Log'!$E:$E,'Price Log'!$B:$B,A{r},'Price Log'!$F:$F,\"sold\",'Price Log'!$A:$A,\">=\"&(TODAY()-90)),\"\")",
        }
        for col, val in inputs.items():
            c = ca[f"{col}{r}"]
            c.value = val
            if col != "A":
                c.font = INPUT
        for col, val in outputs.items():
            c = ca[f"{col}{r}"]
            c.value = val
            c.fill = OUT_FILL
        ca[f"AD{r}"] = "Unknown"
        ca[f"AD{r}"].font = INPUT
        for j, (_, _, fmt) in enumerate(head, 1):
            if fmt:
                ca.cell(i, j).number_format = fmt

    last = len(PLANTS) + 1
    dv_rating = DataValidation(type="list", formula1='"Steady,Rising,Hype,Saturating,Unknown"', allow_blank=True)
    dv_risk = DataValidation(type="list", formula1='"Low,Medium,High"', allow_blank=True)
    ca.add_data_validation(dv_rating)
    ca.add_data_validation(dv_risk)
    dv_rating.add(f"AD2:AD{last + 30}")
    dv_risk.add(f"D2:D{last + 30}")
    ca.freeze_panes = "B2"

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    build()
